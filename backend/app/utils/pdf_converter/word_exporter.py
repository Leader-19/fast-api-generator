from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_text_line(
    document,
    text_items,
) -> None:

    if not text_items:
        return

    text_items.sort(
        key=lambda item: item.x0
    )

    text = " ".join(
        item.text
        for item in text_items
    )

    if not text.strip():
        return

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = (
        Pt(0)
    )

    paragraph.paragraph_format.space_after = (
        Pt(3)
    )

    run = paragraph.add_run(
        text
    )

    first = text_items[0]

    if first.size:

        run.font.size = Pt(
            max(
                8,
                min(
                    first.size,
                    30,
                ),
            )
        )


def add_table(
    document,
    table_data,
) -> None:

    rows = table_data.rows

    if not rows:
        return

    column_count = max(
        len(row)
        for row in rows
    )

    table = document.add_table(
        rows=len(rows),
        cols=column_count,
    )

    table.style = "Table Grid"

    for row_index, row in enumerate(
        rows
    ):

        for column_index in range(
            column_count
        ):

            value = ""

            if column_index < len(row):
                value = row[
                    column_index
                ]

            cell = table.cell(
                row_index,
                column_index,
            )

            cell.text = value or ""

            for paragraph in (
                cell.paragraphs
            ):

                paragraph.paragraph_format.space_before = (
                    Pt(0)
                )

                paragraph.paragraph_format.space_after = (
                    Pt(0)
                )

                for run in paragraph.runs:

                    run.font.size = Pt(
                        9
                    )

                    if row_index == 0:
                        run.bold = True


def export_to_word(
    pages,
    output_path: str,
) -> None:

    if not pages:
        raise ValueError(
            "No content found in PDF."
        )

    document = Document()

    for page_index, page in enumerate(
        pages
    ):

        if page_index > 0:
            document.add_page_break()

        # --------------------------------
        # Page heading
        # --------------------------------

        heading = document.add_paragraph()

        heading.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = heading.add_run(
            f"Page {page.number}"
        )

        run.bold = True

        run.font.size = Pt(12)

        # --------------------------------
        # Tables
        # --------------------------------

        if page.tables:

            sorted_tables = sorted(
                page.tables,
                key=lambda table: (
                    table.top,
                    table.x0,
                ),
            )

            for table in sorted_tables:

                add_table(
                    document,
                    table,
                )

                document.add_paragraph()

        # --------------------------------
        # Text
        # --------------------------------

        if page.texts:

            tolerance = 3.0

            lines = []

            sorted_words = sorted(
                page.texts,
                key=lambda item: (
                    item.top,
                    item.x0,
                ),
            )

            for word in sorted_words:

                found_line = None

                for line in lines:

                    average_top = sum(
                        item.top
                        for item in line
                    ) / len(line)

                    if abs(
                        word.top
                        - average_top
                    ) <= tolerance:

                        found_line = line
                        break

                if found_line is None:

                    lines.append(
                        [word]
                    )

                else:

                    found_line.append(
                        word
                    )

            lines.sort(
                key=lambda line: min(
                    item.top
                    for item in line
                )
            )

            for line in lines:

                add_text_line(
                    document,
                    line,
                )

    document.save(
        output_path
    )