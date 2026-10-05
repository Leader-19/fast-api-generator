from __future__ import annotations

import re

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    Side,
)
from openpyxl.utils import get_column_letter


THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def safe_excel_value(
    value: str,
):
    """
    Convert obvious numbers into Excel numbers.

    Keep other values as text.
    """

    if value is None:
        return ""

    value = str(value).strip()

    if not value:
        return ""

    # Integer
    if re.fullmatch(
        r"-?\d+",
        value,
    ):
        try:
            return int(value)
        except ValueError:
            pass

    # Decimal
    if re.fullmatch(
        r"-?\d+\.\d+",
        value,
    ):
        try:
            return float(value)
        except ValueError:
            pass

    return value


def estimate_column_width(
    value: str,
) -> float:

    if not value:
        return 10

    length = len(
        str(value)
    )

    return min(
        max(length + 2, 10),
        45,
    )


def write_table(
    worksheet,
    table,
    start_row: int,
    start_column: int = 1,
) -> int:

    rows = table.rows

    if not rows:
        return start_row

    max_columns = max(
        len(row)
        for row in rows
    )

    for row_index, row in enumerate(
        rows,
        start=start_row,
    ):

        for column_index in range(
            max_columns
        ):

            value = ""

            if column_index < len(row):
                value = row[
                    column_index
                ]

            cell = worksheet.cell(
                row=row_index,
                column=(
                    start_column
                    + column_index
                ),
                value=safe_excel_value(
                    value
                ),
            )

            cell.border = THIN_BORDER

            cell.alignment = Alignment(
                vertical="center",
                horizontal="left",
                wrap_text=True,
            )

            # First row = header
            if row_index == start_row:

                cell.font = Font(
                    bold=True,
                )

    return start_row + len(rows)


def write_text(
    worksheet,
    page,
    start_row: int,
) -> int:

    if not page.texts:
        return start_row

    # Group words into lines based on
    # their vertical position.
    lines: list[list] = []

    tolerance = 3.0

    sorted_words = sorted(
        page.texts,
        key=lambda item: (
            item.top,
            item.x0,
        ),
    )

    for word in sorted_words:

        found_line = None

        for line in lines:

            average_top = sum(
                item.top
                for item in line
            ) / len(line)

            if abs(
                word.top - average_top
            ) <= tolerance:

                found_line = line
                break

        if found_line is None:

            lines.append(
                [word]
            )

        else:

            found_line.append(
                word
            )

    lines.sort(
        key=lambda line: min(
            item.top
            for item in line
        )
    )

    row = start_row

    for line in lines:

        line.sort(
            key=lambda item: item.x0
        )

        text = " ".join(
            item.text
            for item in line
        )

        if not text.strip():
            continue

        cell = worksheet.cell(
            row=row,
            column=1,
            value=text,
        )

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

        # Use the first text item to
        # estimate font size.
        first = line[0]

        if first.size:

            cell.font = Font(
                size=max(
                    8,
                    min(
                        int(
                            round(
                                first.size
                            )
                        ),
                        30,
                    ),
                ),
            )

        row += 1

    return row


def export_to_excel(
    pages,
    output_path: str,
) -> None:

    if not pages:
        raise ValueError(
            "No content found in PDF."
        )

    workbook = Workbook()

    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )

    for page in pages:

        worksheet = workbook.create_sheet(
            title=f"Page_{page.number}"
        )

        worksheet.sheet_view.showGridLines = (
            False
        )

        current_row = 1

        # --------------------------------
        # Tables
        # --------------------------------

        if page.tables:

            # Sort tables by their
            # original PDF position.
            sorted_tables = sorted(
                page.tables,
                key=lambda table: (
                    table.top,
                    table.x0,
                ),
            )

            for table in sorted_tables:

                current_row = write_table(
                    worksheet,
                    table,
                    current_row,
                )

                current_row += 2

        # --------------------------------
        # Text
        # --------------------------------

        if page.texts:

            current_row = write_text(
                worksheet,
                page,
                current_row,
            )

        # --------------------------------
        # Column widths
        # --------------------------------

        for column_cells in worksheet.columns:

            max_length = 0

            column_index = (
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is None:
                    continue

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    ),
                )

            worksheet.column_dimensions[
                get_column_letter(
                    column_index
                )
            ].width = min(
                max(
                    max_length + 2,
                    10,
                ),
                50,
            )

        # --------------------------------
        # Print settings
        # --------------------------------

        worksheet.page_setup.orientation = (
            "portrait"
        )

        worksheet.page_setup.fitToWidth = 1

        worksheet.sheet_properties.pageSetUpPr.fitToPage = (
            True
        )

        worksheet.freeze_panes = "A1"

    workbook.save(
        output_path
    )