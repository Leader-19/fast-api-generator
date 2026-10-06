from __future__ import annotations

import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sqlalchemy.orm import Session

from app.repositories.reconciliation_repository import (
    ReconciliationRepository,
)
from app.services.reconciliation_service import (
    ReconciliationService,
)


HEADER_FILL = PatternFill(
    start_color="1E293B",
    end_color="1E293B",
    fill_type="solid",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)


def _style_header(
    worksheet,
    headers: list[str],
) -> None:

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=header,
        )

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
        )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = max(len(header) + 4, 14)

    worksheet.freeze_panes = "A2"


def _autosize(
    worksheet,
    row_count: int,
    column_count: int,
) -> None:

    for column_index in range(
        1,
        column_count + 1,
    ):

        longest = len(
            str(
                worksheet.cell(
                    row=1,
                    column=column_index,
                ).value,
            ),
        )

        for row_index in range(
            2,
            min(row_count, 200) + 2,
        ):
            value = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

            if value is not None:
                longest = max(
                    longest,
                    len(str(value)),
                )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(longest + 2, 40)


class ExcelService:

    def __init__(self):
        self.repository = (
            ReconciliationRepository()
        )
        self.reconciliation_service = (
            ReconciliationService()
        )

    def export_run_matches(
        self,
        db: Session,
        run_id: int,
    ) -> str:

        run = self.reconciliation_service.get_run(
            db,
            run_id,
        )

        rows = (
            self.repository.get_matches_with_records(
                db,
                run_id,
            )
        )

        workbook = Workbook()

        # --------------------------------------------------
        # Sheet 1: Summary
        # --------------------------------------------------
        summary = workbook.active
        summary.title = "Summary"

        _style_header(
            summary,
            ["Metric", "Value"],
        )

        summary.append(
            ["Run name", run.name],
        )
        summary.append(["Status", run.status])
        summary.append(
            ["Total records", run.total_records],
        )
        summary.append(
            [
                "Matched",
                run.matched_records,
            ],
        )
        summary.append(
            [
                "Mismatched",
                run.mismatched_records,
            ],
        )
        summary.append(
            [
                "Unmatched",
                run.unmatched_records,
            ],
        )
        summary.append(
            [
                "Total variance",
                run.total_variance,
            ],
        )

        if run.started_at:
            summary.append(
                [
                    "Started at",
                    run.started_at.strftime(
                        "%Y-%m-%d %H:%M:%S",
                    ),
                ],
            )

        if run.completed_at:
            summary.append(
                [
                    "Completed at",
                    run.completed_at.strftime(
                        "%Y-%m-%d %H:%M:%S",
                    ),
                ],
            )

        _autosize(summary, 8, 2)

        # --------------------------------------------------
        # Sheet 2: Matches
        # --------------------------------------------------
        matches_sheet = workbook.create_sheet(
            "Matches",
        )

        _style_header(
            matches_sheet,
            [
                "Match ID",
                "Status",
                "Reference",
                "Partner",
                "Date",
                "Side A charge",
                "Side B charge",
                "Charge variance",
                "Duration variance (s)",
                "Volume variance (KB)",
                "Root cause",
            ],
        )

        for row in rows:

            record_a = row["record_a"]
            record_b = row["record_b"]
            primary = row["primary"]

            matches_sheet.append(
                [
                    row["match"].id,
                    row["match"].status,
                    primary.reference
                    if primary
                    else "",
                    primary.partner
                    if primary
                    else "",
                    primary.record_date.strftime(
                        "%Y-%m-%d %H:%M:%S",
                    )
                    if primary
                    else "",
                    record_a.charge
                    if record_a
                    else None,
                    record_b.charge
                    if record_b
                    else None,
                    round(
                        row[
                            "match"
                        ].charge_variance,
                        2,
                    ),
                    row[
                        "match"
                    ].duration_variance,
                    round(
                        row[
                            "match"
                        ].volume_variance,
                        2,
                    ),
                    row["match"].root_cause
                    or "",
                ],
            )

        _autosize(
            matches_sheet,
            len(rows),
            11,
        )

        # --------------------------------------------------
        # Sheet 3: Root causes
        # --------------------------------------------------
        breakdown = (
            self.repository.get_root_cause_breakdown(
                db,
                run_id,
            )
        )

        causes_sheet = workbook.create_sheet(
            "Root causes",
        )

        _style_header(
            causes_sheet,
            [
                "Root cause",
                "Count",
                "Total variance",
            ],
        )

        for item in breakdown:
            causes_sheet.append(
                [
                    item["root_cause"],
                    item["count"],
                    round(
                        item["total_variance"],
                        2,
                    ),
                ],
            )

        _autosize(
            causes_sheet,
            len(breakdown),
            3,
        )

        output_directory = tempfile.mkdtemp(
            prefix="ra_report_",
        )

        safe_name = "".join(
            character
            if character.isalnum()
            or character in ("- ", "_")
            else "_"
            for character in run.name
        ).strip() or f"run_{run.id}"

        output_path = os.path.join(
            output_directory,
            f"reconciliation_{safe_name}.xlsx",
        )

        workbook.save(output_path)

        return output_path

    def export_summary(
        self,
        db: Session,
    ) -> str:

        data = (
            self.repository.get_run_summary(db)
        )

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "All runs"

        _style_header(
            sheet,
            [
                "Run ID",
                "Name",
                "Status",
                "Total records",
                "Matched",
                "Mismatched",
                "Unmatched",
                "Total variance",
                "Created at",
            ],
        )

        for run in data["runs"]:

            sheet.append(
                [
                    run.id,
                    run.name,
                    run.status,
                    run.total_records,
                    run.matched_records,
                    run.mismatched_records,
                    run.unmatched_records,
                    round(run.total_variance, 2),
                    run.created_at.strftime(
                        "%Y-%m-%d %H:%M:%S",
                    ),
                ],
            )

        _autosize(
            sheet,
            len(data["runs"]),
            9,
        )

        output_directory = tempfile.mkdtemp(
            prefix="ra_summary_",
        )

        output_path = os.path.join(
            output_directory,
            "reconciliation_summary.xlsx",
        )

        workbook.save(output_path)

        return output_path
